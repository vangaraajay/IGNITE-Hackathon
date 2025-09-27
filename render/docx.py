from docx import Document
from docx.shared import Inches
from io import BytesIO

def plan_to_docx(plan: dict) -> bytes:
    doc = Document()
    doc.add_heading("Bring-Up Procedure", 0)

    doc.add_heading("Voltage Rails to Check", level=1)
    rails = plan["procedure"]["rail_table"]
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Item", "Pin", "Net", "Value"
    for r in rails:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = r["item"], r["pin"], r["net"], str(r["value"])

    doc.add_heading("Steps", level=1)
    for s in plan["procedure"]["steps"]:
        doc.add_paragraph(s, style="List Number")

    doc.add_heading("Appendix: Test Datasheet", level=1)
    dt = plan.get("datasheet", [])
    t2 = doc.add_table(rows=1, cols=3)
    h2 = t2.rows[0].cells
    h2[0].text, h2[1].text, h2[2].text = "Section", "Description", "Expected"
    for row in dt:
        r = t2.add_row().cells
        r[0].text, r[1].text, r[2].text = row.get("section",""), row.get("description",""), row.get("expected","")

    if plan.get("assumptions"):
        doc.add_heading("Assumptions", level=1)
        for a in plan["assumptions"]:
            doc.add_paragraph(a, style="List Bullet")

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()
